"""File -> text extraction, one loader per format.

Loaders return ``langchain_core.documents.Document`` objects so the rest of the
ingestion path is format-agnostic. Page-aware formats (PDF) emit one Document
per page, which lets retrieved chunks carry a page number for citation.
"""

from __future__ import annotations

import csv
import io
import json
from collections.abc import Callable, Iterable
from pathlib import Path

from langchain_core.documents import Document

from app.core.exceptions import IngestionError, UnsupportedDocumentError
from app.core.logging import get_logger
from app.utils.text import normalize_whitespace

logger = get_logger(__name__)

Loader = Callable[[Path], list[Document]]


def _base_metadata(path: Path) -> dict[str, object]:
    stat = path.stat()
    return {
        "source": path.name,
        "source_path": str(path),
        "file_type": path.suffix.lower().lstrip("."),
        "size_bytes": stat.st_size,
        "modified_at": int(stat.st_mtime),
    }


def _read_text(path: Path) -> str:
    """Decode with a couple of fallbacks so odd encodings don't kill ingestion."""
    for encoding in ("utf-8", "utf-8-sig", "cp1252", "latin-1"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    raise IngestionError(f"could not decode {path.name} as text")


def load_text(path: Path) -> list[Document]:
    content = normalize_whitespace(_read_text(path))
    if not content:
        return []
    return [Document(page_content=content, metadata=_base_metadata(path))]


def load_pdf(path: Path) -> list[Document]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover
        raise IngestionError("PDF ingestion requires `pip install pypdf`.") from exc

    try:
        reader = PdfReader(str(path))
    except Exception as exc:
        raise IngestionError(f"could not open PDF {path.name}: {exc}") from exc

    base = _base_metadata(path)
    docs: list[Document] = []
    for number, page in enumerate(reader.pages, start=1):
        try:
            text = normalize_whitespace(page.extract_text() or "")
        except Exception as exc:  # a single unreadable page shouldn't abort the file
            logger.warning("skipping page %d of %s: %s", number, path.name, exc)
            continue
        if text:
            docs.append(
                Document(page_content=text, metadata={**base, "page": number, "pages": len(reader.pages)})
            )
    if not docs:
        raise IngestionError(
            f"{path.name} produced no extractable text -- it may be a scanned "
            "PDF that needs OCR."
        )
    return docs


def load_docx(path: Path) -> list[Document]:
    try:
        import docx
    except ImportError as exc:  # pragma: no cover
        raise IngestionError("DOCX ingestion requires `pip install python-docx`.") from exc

    try:
        document = docx.Document(str(path))
    except Exception as exc:
        raise IngestionError(f"could not open DOCX {path.name}: {exc}") from exc

    blocks: list[str] = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))

    content = normalize_whitespace("\n".join(blocks))
    return [Document(page_content=content, metadata=_base_metadata(path))] if content else []


def load_html(path: Path) -> list[Document]:
    raw = _read_text(path)
    try:
        from bs4 import BeautifulSoup

        soup = BeautifulSoup(raw, "html.parser")
        for tag in soup(["script", "style", "nav", "footer", "noscript"]):
            tag.decompose()
        text = soup.get_text(separator="\n")
        title = soup.title.string.strip() if soup.title and soup.title.string else path.stem
    except ImportError:  # pragma: no cover - degrade to a crude strip
        import re

        text = re.sub(r"<[^>]+>", " ", raw)
        title = path.stem

    content = normalize_whitespace(text)
    if not content:
        return []
    return [Document(page_content=content, metadata={**_base_metadata(path), "title": title})]


def load_json(path: Path) -> list[Document]:
    try:
        payload = json.loads(_read_text(path))
    except json.JSONDecodeError as exc:
        raise IngestionError(f"invalid JSON in {path.name}: {exc}") from exc

    base = _base_metadata(path)
    records: Iterable[object] = payload if isinstance(payload, list) else [payload]
    docs: list[Document] = []
    for index, record in enumerate(records):
        if isinstance(record, dict):
            text = "\n".join(f"{key}: {value}" for key, value in record.items())
        else:
            text = str(record)
        text = normalize_whitespace(text)
        if text:
            docs.append(Document(page_content=text, metadata={**base, "record": index}))
    return docs


def load_csv(path: Path) -> list[Document]:
    raw = _read_text(path)
    try:
        dialect = csv.Sniffer().sniff(raw[:4096])
    except csv.Error:
        dialect = csv.excel  # type: ignore[assignment]

    reader = csv.DictReader(io.StringIO(raw), dialect=dialect)
    base = _base_metadata(path)
    docs: list[Document] = []
    for index, row in enumerate(reader):
        text = normalize_whitespace(
            "\n".join(f"{k}: {v}" for k, v in row.items() if k and v)
        )
        if text:
            docs.append(Document(page_content=text, metadata={**base, "row": index}))
    return docs


LOADERS: dict[str, Loader] = {
    ".txt": load_text,
    ".text": load_text,
    ".md": load_text,
    ".markdown": load_text,
    ".rst": load_text,
    ".log": load_text,
    ".pdf": load_pdf,
    ".docx": load_docx,
    ".html": load_html,
    ".htm": load_html,
    ".json": load_json,
    ".jsonl": load_json,
    ".csv": load_csv,
    ".tsv": load_csv,
}

SUPPORTED_EXTENSIONS: frozenset[str] = frozenset(LOADERS)


def load_document(path: Path) -> list[Document]:
    """Dispatch to the loader for ``path``'s extension."""
    path = Path(path)
    if not path.is_file():
        raise IngestionError(f"not a file: {path}")
    loader = LOADERS.get(path.suffix.lower())
    if loader is None:
        raise UnsupportedDocumentError(
            f"unsupported file type {path.suffix!r}",
            details={"supported": sorted(SUPPORTED_EXTENSIONS)},
        )
    return loader(path)
